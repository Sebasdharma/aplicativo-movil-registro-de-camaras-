from flask import Flask, jsonify, render_template, request, send_file
from flask_cors import CORS
import base64
import binascii
from io import BytesIO
import os
import re
import cv2
import numpy as np
import uuid
import jwt
from database import (
    init_db, get_db, registrar_asistencia_db, obtener_resumen_asistencia,
    obtener_materias_docente, obtener_dispositivos_iot,
    inscribir_estudiante_en_materia, aplicar_ausencias_programadas,
    obtener_periodos_lectivos,
)
from carchito_voice import CarchitoVoiceAssistant
from face_engine import FaceEngine
import threading
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

app = Flask(__name__)
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0
app.config['MAX_CONTENT_LENGTH'] = 12 * 1024 * 1024
CORS(app)

MICROSOFT_TENANT_ID = os.environ.get(
    "MICROSOFT_TENANT_ID",
    "0a42bec9-732b-45d1-977d-3b8d3ac98c2b",
)
MICROSOFT_CLIENT_ID = os.environ.get(
    "MICROSOFT_CLIENT_ID",
    "6a39d754-3c15-4599-b074-3925b7c7e77e",
)
MICROSOFT_ISSUER = (
    f"https://login.microsoftonline.com/{MICROSOFT_TENANT_ID}/v2.0"
)
MAX_FOTOS_POR_SOLICITUD = 5
MAX_FOTOS_ADICIONALES_POR_ESTUDIANTE = 10
MICROSOFT_JWKS_CLIENT = jwt.PyJWKClient(
    f"https://login.microsoftonline.com/{MICROSOFT_TENANT_ID}"
    "/discovery/v2.0/keys"
)

init_db()
asistente = CarchitoVoiceAssistant()
face_engine = FaceEngine()
camera_thread = None
camera_thread_lock = threading.Lock()
tablet_cameras = {}
tablet_cameras_lock = threading.Lock()
TABLET_CAMERA_TIMEOUT_SECONDS = 12


def verificar_token_microsoft(id_token):
    signing_key = MICROSOFT_JWKS_CLIENT.get_signing_key_from_jwt(id_token)
    return jwt.decode(
        id_token,
        signing_key.key,
        algorithms=["RS256"],
        audience=MICROSOFT_CLIENT_ID,
        issuer=MICROSOFT_ISSUER,
        options={"require": ["exp", "iss", "aud", "tid", "oid"]},
    )


@app.route("/api/auth/admin", methods=["POST"])
def api_auth_admin():
    conn = get_db()
    try:
        docente = conn.execute(
            """
            SELECT id, email_institucional, nombre_completo, cedula, facultad, carrera
            FROM docentes
            WHERE email_institucional = ?
            """,
            ("admin@upec.edu.ec",),
        ).fetchone()
    finally:
        conn.close()

    if docente is None:
        return jsonify({
            "status": "error",
            "mensaje": "No está configurada la cuenta de administrador.",
        }), 503

    return jsonify({
        "status": "ok",
        "provider": "admin_development",
        "docente": dict(docente),
    })


@app.route("/api/auth/microsoft", methods=["POST"])
def api_auth_microsoft():
    data = request.get_json(silent=True)
    id_token = data.get("id_token") if isinstance(data, dict) else None
    if not isinstance(id_token, str) or not id_token:
        return jsonify({
            "status": "error",
            "mensaje": "Microsoft no envió un token de identidad válido.",
        }), 400

    try:
        claims = verificar_token_microsoft(id_token)
    except jwt.PyJWKClientError:
        app.logger.exception("No se pudo obtener la firma pública de Microsoft.")
        return jsonify({
            "status": "error",
            "mensaje": "No se pudo verificar la respuesta de Microsoft. Intenta de nuevo.",
        }), 503
    except jwt.PyJWTError as error:
        app.logger.warning(
            "Token de Microsoft rechazado (%s).",
            type(error).__name__,
        )
        return jsonify({
            "status": "error",
            "mensaje": "Microsoft devolvió un token inválido o vencido.",
        }), 401

    if claims.get("tid") != MICROSOFT_TENANT_ID:
        return jsonify({
            "status": "error",
            "mensaje": "La cuenta no pertenece al directorio institucional configurado.",
        }), 403

    email = claims.get("preferred_username") or claims.get("email")
    if not isinstance(email, str) or not email:
        return jsonify({
            "status": "error",
            "mensaje": "Microsoft no proporcionó un correo para esta cuenta.",
        }), 403

    conn = get_db()
    try:
        docente = conn.execute(
            """
            SELECT id, email_institucional, nombre_completo, cedula, facultad, carrera
            FROM docentes
            WHERE LOWER(email_institucional) = LOWER(?)
            """,
            (email,),
        ).fetchone()
    finally:
        conn.close()

    if docente is None:
        return jsonify({
            "status": "error",
            "mensaje": (
                "Tu cuenta Microsoft se autenticó, pero aún no está registrada "
                "como docente de AsisTech. Para las pruebas actuales, ingresa "
                "con el acceso temporal de Administrador."
            ),
        }), 403

    return jsonify({
        "status": "ok",
        "provider": "microsoft",
        "docente": dict(docente),
    })


@app.route("/api/materias", methods=["GET"])
def api_materias():
    docente_id = request.args.get("docente_id", 1, type=int)
    return jsonify(obtener_materias_docente(docente_id))

@app.route("/api/asistencia/resumen", methods=["GET"])
def api_resumen():
    materia_id = request.args.get("materia_id", 1, type=int)
    fecha = request.args.get("fecha")
    try:
        return jsonify(obtener_resumen_asistencia(materia_id, fecha))
    except ValueError as error:
        return jsonify({"status": "error", "mensaje": str(error)}), 400
    except Exception:
        app.logger.exception("No se pudo consultar la asistencia por periodos.")
        return jsonify({
            "status": "error",
            "mensaje": "No se pudo consultar la asistencia. Intenta de nuevo.",
        }), 500


@app.route("/api/asistencia/reporte/excel", methods=["GET"])
def api_reporte_excel():
    materia_id = request.args.get("materia_id", type=int)
    docente_id = request.args.get("docente_id", type=int)
    fecha = request.args.get("fecha", datetime.now().strftime("%Y-%m-%d"))
    if not materia_id or not docente_id:
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona la materia y el docente del reporte.",
        }), 400

    materias_docente = obtener_materias_docente(docente_id)
    materia = next(
        (item for item in materias_docente if item["id"] == materia_id),
        None,
    )
    if materia is None:
        return jsonify({
            "status": "error",
            "mensaje": "La materia no pertenece a la cuenta seleccionada.",
        }), 404

    try:
        registros = obtener_resumen_asistencia(materia_id, fecha)
        periodos = (
            registros[0]["periodos"]
            if registros
            else obtener_periodos_lectivos(materia["horario"], fecha)
        )
        total_con_falta = sum(
            any(periodo["estado"] == "Ausente" for periodo in estudiante["periodos"])
            for estudiante in registros
        )

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Asistencia diaria"
        hoja.cell(1, 1, "Reporte diario de asistencia")
        hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + len(periodos))
        hoja.cell(1, 1).font = Font(bold=True, size=14, color="FFFFFF")
        hoja.cell(1, 1).fill = PatternFill("solid", fgColor="006B3F")
        hoja.cell(2, 1, "Materia")
        hoja.cell(2, 2, materia["nombre_materia"])
        hoja.cell(3, 1, "Código")
        hoja.cell(3, 2, materia["codigo_materia"])
        hoja.cell(4, 1, "Fecha")
        hoja.cell(4, 2, fecha)
        hoja.cell(5, 1, "Total de estudiantes")
        hoja.cell(5, 2, len(registros))
        hoja.cell(5, 3, "Estudiantes con falta (únicos)")
        hoja.cell(5, 4, total_con_falta)

        encabezados = [
            "Código estudiantil",
            "Estudiante",
            "Nivel / paralelo",
            *[
                f"{periodo['inicio']}-{periodo['fin']}"
                for periodo in periodos
            ],
            "Falta del día",
        ]
        fila_encabezado = 7
        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(fila_encabezado, columna, encabezado)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="006B3F")
            celda.alignment = Alignment(horizontal="center", wrap_text=True)

        for numero_fila, estudiante in enumerate(registros, start=fila_encabezado + 1):
            detalle = {
                periodo["indice"]: periodo
                for periodo in estudiante["periodos"]
            }
            tiene_falta = any(
                periodo["estado"] == "Ausente"
                for periodo in estudiante["periodos"]
            )
            valores = [
                estudiante["codigo_estudiantil"],
                estudiante["nombre_completo"],
                estudiante["nivel_paralelo"],
                *[
                    (
                        f"{detalle[periodo['indice']]['estado']} "
                        f"{detalle[periodo['indice']]['hora']}"
                        if detalle[periodo["indice"]]["hora"] != "--:--:--"
                        else detalle[periodo["indice"]]["estado"]
                    )
                    for periodo in periodos
                ],
                "Sí" if tiene_falta else "No",
            ]
            for columna, valor in enumerate(valores, start=1):
                celda = hoja.cell(numero_fila, columna, valor)
                if valor == "Sí":
                    celda.fill = PatternFill("solid", fgColor="FCE4D6")
                elif valor == "Presente":
                    celda.fill = PatternFill("solid", fgColor="E2F0D9")

        hoja.freeze_panes = "A8"
        hoja.auto_filter.ref = (
            f"A7:{hoja.cell(fila_encabezado, len(encabezados)).column_letter}"
            f"{max(fila_encabezado, hoja.max_row)}"
        )
        anchos = [20, 34, 22] + [18] * len(periodos) + [16]
        for columna, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[
                hoja.cell(fila_encabezado, columna).column_letter
            ].width = ancho

        archivo = BytesIO()
        libro.save(archivo)
        archivo.seek(0)
        codigo_seguro = re.sub(
            r"[^A-Za-z0-9_-]+", "-", materia["codigo_materia"]
        ).strip("-")
        return send_file(
            archivo,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name=f"asistencia-{codigo_seguro}-{fecha}.xlsx",
        )
    except ValueError as error:
        return jsonify({"status": "error", "mensaje": str(error)}), 400
    except Exception:
        app.logger.exception("No se pudo generar el reporte diario de Excel.")
        return jsonify({
            "status": "error",
            "mensaje": "No se pudo generar el reporte. Intenta de nuevo.",
        }), 500


@app.route("/api/asistencia/manual", methods=["POST"])
def api_manual():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({
            "status": "error",
            "mensaje": "Envía los datos de asistencia como un objeto JSON.",
        }), 400

    est_id = data.get("estudiante_id")
    mat_id = data.get("materia_id", 1)
    estado = data.get("estado", "Presente")
    metodo = data.get("metodo", "Manual-Docente")
    dispositivo_id = data.get("dispositivo_id")
    periodo_indice = data.get("periodo_indice")
    if isinstance(est_id, bool) or not isinstance(est_id, int) or est_id < 1:
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona un estudiante válido.",
        }), 400
    if isinstance(mat_id, bool) or not isinstance(mat_id, int) or mat_id < 1:
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona una materia válida.",
        }), 400
    if not isinstance(estado, str) or estado not in {
        "Presente",
        "Ausente",
        "Atrasado",
        "Justificado",
    }:
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona un estado de asistencia válido.",
        }), 400
    if not isinstance(metodo, str) or not metodo.strip() or len(metodo) > 40:
        return jsonify({
            "status": "error",
            "mensaje": "El método de registro no es válido.",
        }), 400
    if dispositivo_id is not None and (
        isinstance(dispositivo_id, bool)
        or not isinstance(dispositivo_id, int)
        or dispositivo_id < 1
    ):
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona un dispositivo válido.",
        }), 400
    if periodo_indice is not None and (
        isinstance(periodo_indice, bool)
        or not isinstance(periodo_indice, int)
        or periodo_indice < 1
    ):
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona una hora de clase válida.",
        }), 400

    try:
        res = registrar_asistencia_db(
            est_id,
            mat_id,
            estado=estado,
            metodo=metodo,
            dispositivo_id=dispositivo_id,
            periodo_indice=periodo_indice,
        )
        return jsonify(res)
    except ValueError as error:
        return jsonify({"status": "error", "mensaje": str(error)}), 400
    except Exception:
        app.logger.exception("No se pudo registrar la asistencia manual.")
        return jsonify({
            "status": "error",
            "mensaje": "No se pudo guardar la asistencia. Intenta de nuevo.",
        }), 500

@app.route("/api/dispositivos", methods=["GET"])
def api_dispositivos():
    return jsonify(obtener_dispositivos_iot())

@app.route("/api/carchito/comando", methods=["POST"])
def api_comando():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({
            "accion": "error",
            "mensaje": "Envía el comando como un objeto JSON.",
        }), 400

    comando = data.get("comando")
    mat_id = data.get("materia_id", 1)
    if not isinstance(comando, str) or not comando.strip():
        return jsonify({
            "accion": "error",
            "mensaje": "Escribe o dicta un comando para Carchito.",
        }), 400
    if len(comando) > 500:
        return jsonify({
            "accion": "error",
            "mensaje": "El comando es demasiado largo.",
        }), 400
    if isinstance(mat_id, bool) or not isinstance(mat_id, int) or mat_id < 1:
        return jsonify({
            "accion": "error",
            "mensaje": "Selecciona una materia válida antes de hablar con Carchito.",
        }), 400

    res = asistente.procesar_comando(
        comando,
        mat_id,
        muestras_faciales=len(face_engine.label_to_student),
    )
    return jsonify(res)

@app.route("/api/camara/iniciar", methods=["POST"])
def api_camara():
    global camera_thread
    data = request.json or {}
    mat_id = data.get("materia_id", 1)
    if isinstance(mat_id, bool) or not isinstance(mat_id, int) or mat_id < 1:
        return jsonify({
            "status": "error",
            "mensaje": "Selecciona una materia válida antes de iniciar el escaneo.",
        }), 400

    with camera_thread_lock:
        if camera_thread is not None and camera_thread.is_alive():
            return jsonify({
                "status": "ok",
                "mensaje": "El escaneo OpenCV ya está activo.",
            })
        face_engine._camera_stop_event.clear()
        camera_thread = threading.Thread(
            target=face_engine.ejecutar_camara_en_vivo,
            args=(mat_id,),
            daemon=True,
        )
        camera_thread.start()
    return jsonify({
        "status": "ok",
        "mensaje": "Cámara de reconocimiento facial iniciada.",
    })


@app.route("/api/camara/detener", methods=["POST"])
def api_detener_camara():
    face_engine.detener_camara_en_vivo()
    return jsonify({
        "status": "ok",
        "mensaje": "Se solicitó terminar el escaneo OpenCV.",
    })


@app.route("/api/camara/tablet/dispositivos", methods=["GET"])
def api_tablet_cameras():
    now = time.monotonic()
    with tablet_cameras_lock:
        connected = [
            {
                "tablet_id": tablet_id,
                "nombre": camera["nombre"],
                "estado": "Transmitiendo" if camera["activa"] else "Conectada",
                "resultado": camera["resultado"],
            }
            for tablet_id, camera in tablet_cameras.items()
            if now - camera["ultima_actividad"] <= TABLET_CAMERA_TIMEOUT_SECONDS
        ]
    return jsonify(connected)


@app.route("/tablet-camera", methods=["GET"])
def tablet_camera_page():
    return render_template("tablet_camera.html")


@app.route("/api/camara/tablet/heartbeat", methods=["POST"])
def api_tablet_camera_heartbeat():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "mensaje": "Envía datos válidos."}), 400

    tablet_id = data.get("tablet_id")
    nombre = data.get("nombre", "Tablet Android")
    activa = data.get("activa", False)
    if not isinstance(tablet_id, str) or not re.fullmatch(
        r"[A-Za-z0-9_-]{1,64}", tablet_id
    ):
        return jsonify({
            "status": "error",
            "mensaje": "El identificador de la tablet no es válido.",
        }), 400
    if not isinstance(nombre, str) or not nombre.strip() or len(nombre) > 80:
        return jsonify({"status": "error", "mensaje": "El nombre no es válido."}), 400
    if not isinstance(activa, bool):
        return jsonify({"status": "error", "mensaje": "El estado no es válido."}), 400

    if data.get("habilitada") is False:
        with tablet_cameras_lock:
            tablet_cameras.pop(tablet_id, None)
        return jsonify({"status": "ok", "command": None})

    now = time.monotonic()
    with tablet_cameras_lock:
        camera = tablet_cameras.get(tablet_id)
        if camera is None:
            camera = {
                "nombre": nombre.strip(),
                "ultima_actividad": now,
                "activa": activa,
                "command": None,
                "resultado": None,
            }
            tablet_cameras[tablet_id] = camera
        camera["nombre"] = nombre.strip()
        camera["ultima_actividad"] = now
        camera["activa"] = activa
        command = camera["command"]
        camera["command"] = None

    return jsonify({"status": "ok", "command": command})


@app.route("/api/camara/tablet/resultado", methods=["POST"])
def api_tablet_camera_result():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "mensaje": "Envía datos válidos."}), 400

    tablet_id = data.get("tablet_id")
    event_id = data.get("evento_id")
    message = data.get("mensaje")
    result_type = data.get("tipo")
    if not isinstance(tablet_id, str) or not re.fullmatch(
        r"[A-Za-z0-9_-]{1,64}", tablet_id
    ):
        return jsonify({"status": "error", "mensaje": "La tablet no es válida."}), 400
    if not isinstance(event_id, str) or not re.fullmatch(
        r"[A-Za-z0-9_-]{1,64}", event_id
    ):
        return jsonify({"status": "error", "mensaje": "El evento no es válido."}), 400
    if not isinstance(message, str) or not message.strip() or len(message) > 500:
        return jsonify({"status": "error", "mensaje": "El resultado no es válido."}), 400
    if not isinstance(result_type, str) or result_type not in {
        "reconocido",
        "fuera_horario",
        "desconocido",
    }:
        return jsonify({"status": "error", "mensaje": "El tipo de resultado no es válido."}), 400

    now = time.monotonic()
    with tablet_cameras_lock:
        camera = tablet_cameras.get(tablet_id)
        if (
            camera is None
            or now - camera["ultima_actividad"] > TABLET_CAMERA_TIMEOUT_SECONDS
        ):
            return jsonify({
                "status": "error",
                "mensaje": "La tablet no está conectada.",
            }), 404
        camera["resultado"] = {
            "evento_id": event_id,
            "mensaje": message.strip(),
            "tipo": result_type,
        }

    return jsonify({"status": "ok"})


@app.route("/api/camara/tablet/comando", methods=["POST"])
def api_tablet_camera_command():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "mensaje": "Envía datos válidos."}), 400

    tablet_id = data.get("tablet_id")
    action = data.get("action")
    materia_id = data.get("materia_id")
    if not isinstance(tablet_id, str) or not re.fullmatch(
        r"[A-Za-z0-9_-]{1,64}", tablet_id
    ):
        return jsonify({"status": "error", "mensaje": "Selecciona una tablet válida."}), 400
    if action not in ("iniciar", "detener"):
        return jsonify({"status": "error", "mensaje": "La orden no es válida."}), 400
    if action == "iniciar" and (
        isinstance(materia_id, bool)
        or not isinstance(materia_id, int)
        or materia_id < 1
    ):
        return jsonify({"status": "error", "mensaje": "Selecciona una materia válida."}), 400

    now = time.monotonic()
    with tablet_cameras_lock:
        camera = tablet_cameras.get(tablet_id)
        if (
            camera is None
            or now - camera["ultima_actividad"] > TABLET_CAMERA_TIMEOUT_SECONDS
        ):
            return jsonify({
                "status": "error",
                "mensaje": (
                    "La tablet no está conectada. Abre la página de cámara "
                    "en Chrome y concede permiso para usar la cámara."
                ),
            }), 404
        camera["command"] = {
            "action": action,
            "materia_id": materia_id,
        }

    return jsonify({
        "status": "ok",
        "mensaje": (
            "Se envió la orden a la tablet."
            if action == "iniciar"
            else "Se solicitó detener la cámara de la tablet."
        ),
    })


@app.route("/api/camara/telefono/escanear", methods=["POST"])
def api_escanear_camara_telefono():
    data = request.json or {}
    foto = data.get("foto")
    materia_id = data.get("materia_id", 1)
    rotacion = data.get("rotacion", 0)
    if rotacion not in (0, 90, 180, 270):
        return jsonify({"status": "error", "mensaje": "La rotación de cámara no es válida."}), 400
    if not isinstance(foto, str) or not foto:
        return jsonify({"status": "error", "mensaje": "Se requiere una foto capturada por la cámara."}), 400
    try:
        image_bytes = base64.b64decode(foto, validate=True)
    except (ValueError, binascii.Error):
        return jsonify({"status": "error", "mensaje": "La foto enviada no tiene un formato válido."}), 400
    if len(image_bytes) > 6 * 1024 * 1024:
        return jsonify({"status": "error", "mensaje": "La foto supera el tamaño máximo permitido de 6 MB."}), 413

    result, status_code = face_engine.reconocer_imagen(
        image_bytes,
        materia_id,
        rotacion,
        session_key=f"{request.remote_addr or 'unknown'}:{materia_id}",
    )
    return jsonify(result), status_code

@app.route("/api/estudiantes/enrolar", methods=["POST"])
def api_enrolar():
    data = request.get_json(silent=True)
    if not isinstance(data, dict):
        return jsonify({"status": "error", "mensaje": "Envía los datos del estudiante como JSON."}), 400
    codigo = data.get("codigo")
    nombre = data.get("nombre")
    paralelo = data.get("paralelo", "7mo Semestre - Paralelo AM")
    email = data.get("email")
    materia_id = data.get("materia_id")
    fotos = data.get("fotos")
    if fotos is None and data.get("foto") is not None:
        fotos = [data["foto"]]

    if not codigo or not nombre:
        return jsonify({"status": "error", "mensaje": "Código y nombre requeridos"}), 400
    if not isinstance(codigo, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,40}", codigo):
        return jsonify({"status": "error", "mensaje": "El código estudiantil contiene caracteres no válidos."}), 400
    if not isinstance(nombre, str) or not nombre.strip() or len(nombre) > 160:
        return jsonify({"status": "error", "mensaje": "Ingresa un nombre válido para el estudiante."}), 400
    if isinstance(materia_id, bool) or not isinstance(materia_id, int) or materia_id < 1:
        return jsonify({"status": "error", "mensaje": "Selecciona una materia válida."}), 400
    images = []
    if fotos is not None:
        if (
            not isinstance(fotos, list)
            or not fotos
            or len(fotos) > MAX_FOTOS_POR_SOLICITUD
            or any(not isinstance(foto, str) or not foto for foto in fotos)
        ):
            return jsonify({
                "status": "error",
                "mensaje": (
                    "Envía de 1 a "
                    f"{MAX_FOTOS_POR_SOLICITUD} fotos faciales por solicitud."
                ),
            }), 400
        if data.get("consentimiento_biometrico") is not True:
            return jsonify({"status": "error", "mensaje": "Se requiere confirmar la autorización antes de guardar las fotos faciales."}), 400
        total_image_bytes = 0
        for foto in fotos:
            try:
                image_bytes = base64.b64decode(foto, validate=True)
            except (ValueError, binascii.Error):
                return jsonify({"status": "error", "mensaje": "Una de las fotos enviadas no tiene un formato válido."}), 400
            total_image_bytes += len(image_bytes)
            if len(image_bytes) > 6 * 1024 * 1024:
                return jsonify({"status": "error", "mensaje": "Cada foto debe pesar menos de 6 MB."}), 413
            if total_image_bytes > 8 * 1024 * 1024:
                return jsonify({"status": "error", "mensaje": "El conjunto de fotos supera el límite total de 8 MB."}), 413
            valido, mensaje = face_engine.validar_foto_enrolamiento(image_bytes)
            if not valido:
                return jsonify({"status": "error", "mensaje": mensaje}), 400
            image = cv2.imdecode(
                np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR
            )
            if image is None:
                return jsonify({"status": "error", "mensaje": "No se pudo leer una de las fotos enviadas."}), 400
            encoded_ok, encoded_image = cv2.imencode(".jpg", image)
            if not encoded_ok:
                return jsonify({"status": "error", "mensaje": "No se pudo preparar una de las fotos faciales."}), 400
            stored_image_bytes = encoded_image.tobytes()
            valido, mensaje = face_engine.validar_foto_enrolamiento(
                stored_image_bytes
            )
            if not valido:
                return jsonify({"status": "error", "mensaje": mensaje}), 400
            images.append(stored_image_bytes)

    conn = get_db()
    cursor = conn.cursor()
    saved_photo_paths = []
    try:
        if cursor.execute("SELECT 1 FROM materias WHERE id = ?", (materia_id,)).fetchone() is None:
            return jsonify({"status": "error", "mensaje": "La materia seleccionada no existe."}), 404

        cursor.execute("SELECT id FROM estudiantes WHERE codigo_estudiantil = ?", (codigo,))
        existente = cursor.fetchone()
        if existente:
            estudiante_id = existente["id"]
            cursor.execute("""
            UPDATE estudiantes
            SET nombre_completo = ?, nivel_paralelo = ?, email = ?, activo = 1
            WHERE id = ?
            """, (nombre.strip(), paralelo, email, estudiante_id))
        else:
            cursor.execute("""
            INSERT INTO estudiantes (codigo_estudiantil, nombre_completo, nivel_paralelo, email, foto_path)
            VALUES (?, ?, ?, ?, ?)
            """, (codigo, nombre.strip(), paralelo, email, None))
            estudiante_id = cursor.lastrowid

        if images:
            fotos_adicionales_anteriores = cursor.execute(
                """
                SELECT COUNT(*) AS total
                FROM fotos_faciales_estudiantes
                WHERE estudiante_id = ?
                """,
                (estudiante_id,),
            ).fetchone()["total"]
            if (
                fotos_adicionales_anteriores + len(images)
                > MAX_FOTOS_ADICIONALES_POR_ESTUDIANTE
            ):
                return jsonify({
                    "status": "error",
                    "mensaje": (
                        f"Este estudiante ya tiene {fotos_adicionales_anteriores} "
                        f"fotos adicionales; puede agregar hasta "
                        f"{MAX_FOTOS_ADICIONALES_POR_ESTUDIANTE} sin contar "
                        "la foto de referencia anterior."
                    ),
                }), 409

        inscribir_estudiante_en_materia(cursor, estudiante_id, materia_id)
        for image_bytes in images:
            foto_filename = f"{codigo}-{uuid.uuid4().hex}.jpg"
            foto_path = os.path.join(face_engine.dataset_dir, foto_filename)
            with open(foto_path, "wb") as photo_file:
                photo_file.write(image_bytes)
            saved_photo_paths.append(foto_path)
            cursor.execute(
                """
                INSERT INTO fotos_faciales_estudiantes (estudiante_id, foto_path)
                VALUES (?, ?)
                """,
                (estudiante_id, foto_filename),
            )
        conn.commit()
        face_engine.cargar_dataset_fotos()
        mensaje = (
            f"Estudiante inscrito y {len(images)} foto(s) facial(es) enrolada(s). "
            "Las referencias anteriores se conservaron."
            if images
            else "Estudiante inscrito en la materia. La referencia facial se puede agregar después."
        )
        return jsonify({"status": "ok", "mensaje": mensaje})
    except Exception:
        conn.rollback()
        for foto_path in saved_photo_paths:
            if os.path.exists(foto_path):
                os.remove(foto_path)
        app.logger.exception("No se pudo guardar el enrolamiento facial.")
        return jsonify({"status": "error", "mensaje": "No se pudo completar el enrolamiento."}), 500
    finally:
        conn.close()

if __name__ == "__main__":
    tls_cert = os.environ.get("ASISTECH_TLS_CERT")
    tls_key = os.environ.get("ASISTECH_TLS_KEY")
    if bool(tls_cert) != bool(tls_key):
        raise RuntimeError(
            "Configura ASISTECH_TLS_CERT y ASISTECH_TLS_KEY juntos para HTTPS."
        )
    ssl_context = (tls_cert, tls_key) if tls_cert and tls_key else None
    protocol = "https" if ssl_context else "http"
    print(f"Iniciando API AsisTech UPEC en {protocol}://0.0.0.0:5000 ...")

    def actualizar_ausencias_programadas():
        while True:
            try:
                aplicar_ausencias_programadas()
            except Exception:
                app.logger.exception(
                    "No se pudieron actualizar las faltas del horario."
                )
            threading.Event().wait(10)

    threading.Thread(
        target=actualizar_ausencias_programadas,
        name="actualizador-asistencias",
        daemon=True,
    ).start()
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False,
        ssl_context=ssl_context,
    )
