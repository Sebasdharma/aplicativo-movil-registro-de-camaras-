import os
import tempfile
import unittest
from datetime import datetime, timedelta
from io import BytesIO
from unittest.mock import patch

import database
import jwt
from openpyxl import load_workbook

_server_import_dir = tempfile.TemporaryDirectory()
_server_import_db_patch = patch.object(
    database,
    "DB_PATH",
    os.path.join(_server_import_dir.name, "server-import.db"),
)
_server_import_db_patch.start()
try:
    import server
finally:
    _server_import_db_patch.stop()
    _server_import_dir.cleanup()

class AdminAuthTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path_patch = patch.object(
            database,
            "DB_PATH",
            os.path.join(self.temp_dir.name, "admin-test.db"),
        )
        self.db_path_patch.start()
        self.addCleanup(self.db_path_patch.stop)
        self.addCleanup(self.temp_dir.cleanup)
        server.init_db()
        with server.tablet_cameras_lock:
            server.tablet_cameras.clear()
        self.client = server.app.test_client()

    def test_admin_login_returns_generic_administrator_profile(self):
        response = self.client.post("/api/auth/admin")

        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["provider"], "admin_development")
        self.assertEqual(result["docente"]["nombre_completo"], "Administrador")
        self.assertEqual(result["docente"]["email_institucional"], "admin@upec.edu.ec")

    def test_unconfigured_microsoft_login_is_not_accepted_as_admin(self):
        response = self.client.post("/api/auth/microsoft", json={
            "email": "someone@upec.edu.ec",
            "provider": "microsoft_azure_ad",
        })

        self.assertEqual(response.status_code, 400)

    def test_microsoft_login_requires_an_identity_token(self):
        with patch.object(server, "verificar_token_microsoft") as verify_token:
            response = self.client.post("/api/auth/microsoft", json={
                "email": "admin@upec.edu.ec",
            })

        self.assertEqual(response.status_code, 400)
        verify_token.assert_not_called()

    def test_daily_excel_report_counts_one_absent_student_and_includes_hour_details(self):
        conn = database.get_db()
        conn.execute(
            """
            INSERT INTO estudiantes
                (codigo_estudiantil, nombre_completo, nivel_paralelo)
            VALUES ('XLSX-001', 'Estudiante de Prueba', 'Séptimo A')
            """
        )
        estudiante_id = conn.execute(
            "SELECT id FROM estudiantes WHERE codigo_estudiantil = 'XLSX-001'"
        ).fetchone()["id"]
        database.inscribir_estudiante_en_materia(
            conn.cursor(), estudiante_id, 1
        )
        materia = conn.execute(
            "SELECT horario FROM materias WHERE id = 1"
        ).fetchone()
        fecha = None
        periodos = []
        for offset in range(8):
            candidate = (datetime.now().date() + timedelta(days=offset)).isoformat()
            candidate_periodos = database.obtener_periodos_lectivos(
                materia["horario"], candidate
            )
            if len(candidate_periodos) >= 2:
                fecha = candidate
                periodos = candidate_periodos
                break
        self.assertIsNotNone(fecha)

        conn.executemany(
            """
            INSERT INTO asistencia_periodos (
                estudiante_id, materia_id, fecha, periodo_indice,
                bloque_indice, inicio, fin, estado, hora, metodo_registro
            ) VALUES (?, 1, ?, ?, ?, ?, ?, ?, ?, 'Manual-Docente')
            """,
            [
                (
                    estudiante_id,
                    fecha,
                    periodo["indice"],
                    periodo["bloque"],
                    periodo["inicio"],
                    periodo["fin"],
                    "Presente" if periodo["indice"] == 1 else "Ausente",
                    "10:05:00" if periodo["indice"] == 1 else "--:--:--",
                )
                for periodo in periodos
            ],
        )
        conn.commit()
        conn.close()

        conn = database.get_db()
        docente_id = conn.execute(
            "SELECT id FROM docentes WHERE email_institucional = 'admin@upec.edu.ec'"
        ).fetchone()["id"]
        conn.close()
        response = self.client.get(
            "/api/asistencia/reporte/excel",
            query_string={
                "materia_id": 1,
                "docente_id": docente_id,
                "fecha": fecha,
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        sheet = workbook["Asistencia diaria"]
        self.assertEqual(sheet["B5"].value, 1)
        self.assertEqual(sheet["D5"].value, 1)
        self.assertEqual(sheet["D8"].value, "Presente 10:05:00")
        self.assertEqual(sheet.cell(8, sheet.max_column).value, "Sí")
        workbook.close()

    def test_microsoft_login_rejects_invalid_identity_token(self):
        with patch.object(
            server,
            "verificar_token_microsoft",
            side_effect=jwt.InvalidTokenError("invalid signature"),
        ):
            response = self.client.post("/api/auth/microsoft", json={
                "id_token": "invalid-token",
            })

        self.assertEqual(response.status_code, 401)

    def test_valid_student_account_is_not_granted_teacher_access(self):
        claims = {
            "tid": server.MICROSOFT_TENANT_ID,
            "oid": "student-object-id",
            "preferred_username": "student@upec.edu.ec",
        }
        with patch.object(server, "verificar_token_microsoft", return_value=claims):
            response = self.client.post("/api/auth/microsoft", json={
                "id_token": "signed-student-token",
            })

        self.assertEqual(response.status_code, 403)
        self.assertIn("no está registrada como docente", response.get_json()["mensaje"])

    def test_valid_registered_teacher_can_sign_in_with_microsoft(self):
        conn = database.get_db()
        conn.execute("""
            INSERT INTO docentes (
                email_institucional, nombre_completo, facultad, carrera
            ) VALUES (?, ?, ?, ?)
        """, (
            "teacher@upec.edu.ec",
            "Docente de prueba",
            "FIACA",
            "Computación",
        ))
        conn.commit()
        conn.close()

        claims = {
            "tid": server.MICROSOFT_TENANT_ID,
            "oid": "teacher-object-id",
            "preferred_username": "teacher@upec.edu.ec",
        }
        with patch.object(server, "verificar_token_microsoft", return_value=claims):
            response = self.client.post("/api/auth/microsoft", json={
                "id_token": "signed-teacher-token",
            })

        self.assertEqual(response.status_code, 200)
        result = response.get_json()
        self.assertEqual(result["provider"], "microsoft")
        self.assertEqual(result["docente"]["nombre_completo"], "Docente de prueba")


class ManualAttendanceApiTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path_patch = patch.object(
            database,
            "DB_PATH",
            os.path.join(self.temp_dir.name, "attendance-api-test.db"),
        )
        self.db_path_patch.start()
        self.addCleanup(self.db_path_patch.stop)
        self.addCleanup(self.temp_dir.cleanup)
        server.init_db()
        self.client = server.app.test_client()

        conn = database.get_db()
        cursor = conn.execute(
            """
            INSERT INTO estudiantes (
                codigo_estudiantil, nombre_completo, nivel_paralelo
            ) VALUES (?, ?, ?)
            """,
            ("API-001", "Estudiante API", "Séptimo A"),
        )
        self.student_id = cursor.lastrowid
        database.inscribir_estudiante_en_materia(conn.cursor(), self.student_id, 1)
        conn.commit()
        conn.close()

    def test_manual_attendance_api_persists_attendance_and_returns_aggregate_status(self):
        with patch.object(database, "datetime", wraps=datetime) as clock:
            clock.now.return_value = datetime(2026, 10, 27, 10, 5, 0)
            response = self.client.post("/api/asistencia/manual", json={
                "estudiante_id": self.student_id,
                "materia_id": 1,
                "estado": "Presente",
                "periodo_indice": 1,
            })
            summary = database.obtener_resumen_asistencia(
                1,
                "2026-10-27",
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["status"], "ok")
        self.assertEqual(summary[0]["estado"], "Presente")
        self.assertEqual(summary[0]["periodos"][0]["estado"], "Presente")

    def test_manual_attendance_api_rejects_malformed_status(self):
        response = self.client.post("/api/asistencia/manual", json={
            "estudiante_id": self.student_id,
            "materia_id": 1,
            "estado": [],
        })

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["status"], "error")


class TabletCameraApiTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        db_path_patch = patch.object(
            database,
            "DB_PATH",
            os.path.join(self.temp_dir.name, "tablet-test.db"),
        )
        db_path_patch.start()
        self.addCleanup(db_path_patch.stop)
        self.addCleanup(self.temp_dir.cleanup)
        server.init_db()
        with server.tablet_cameras_lock:
            server.tablet_cameras.clear()
        self.client = server.app.test_client()

    def test_browser_camera_page_is_available(self):
        response = self.client.get("/tablet-camera")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Permitir c\xc3\xa1mara y conectar", response.data)
        self.assertIn(b"facingMode: { ideal: 'user' }", response.data)
        self.assertIn(b"speechSynthesis.speak(utterance)", response.data)
        self.assertIn(b"fuera del horario de clases", response.data)
        self.assertIn(b"/api/camara/tablet/resultado", response.data)

    def test_phone_can_discover_tablet_and_deliver_start_command_once(self):
        heartbeat = self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-001",
                "nombre": "Tablet AsisTech",
                "activa": False,
            },
        )
        self.assertEqual(heartbeat.status_code, 200)

        devices = self.client.get("/api/camara/tablet/dispositivos")
        self.assertEqual(devices.status_code, 200)
        self.assertEqual(devices.get_json()[0]["tablet_id"], "tablet-001")

        command = self.client.post(
            "/api/camara/tablet/comando",
            json={
                "tablet_id": "tablet-001",
                "action": "iniciar",
                "materia_id": 1,
            },
        )
        self.assertEqual(command.status_code, 200)

        received = self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-001",
                "nombre": "Tablet AsisTech",
                "activa": True,
            },
        )
        self.assertEqual(
            received.get_json()["command"],
            {"action": "iniciar", "materia_id": 1},
        )
        next_poll = self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-001",
                "nombre": "Tablet AsisTech",
                "activa": True,
            },
        )
        self.assertIsNone(next_poll.get_json()["command"])

    def test_tablet_can_publish_recognition_result_to_phone(self):
        self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-resultado",
                "nombre": "Tablet Chrome",
                "activa": True,
            },
        )

        response = self.client.post(
            "/api/camara/tablet/resultado",
            json={
                "tablet_id": "tablet-resultado",
                "evento_id": "evento-001",
                "tipo": "fuera_horario",
                "mensaje": (
                    "Estudiante reconocido: Prueba. Está fuera del horario "
                    "de clases y no se registró la asistencia."
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        devices = self.client.get("/api/camara/tablet/dispositivos").get_json()
        self.assertEqual(
            devices[0]["resultado"],
            {
                "evento_id": "evento-001",
                "tipo": "fuera_horario",
                "mensaje": (
                    "Estudiante reconocido: Prueba. Está fuera del horario "
                    "de clases y no se registró la asistencia."
                ),
            },
        )

    def test_tablet_result_rejects_unregistered_camera(self):
        response = self.client.post(
            "/api/camara/tablet/resultado",
            json={
                "tablet_id": "tablet-desconocida",
                "evento_id": "evento-001",
                "tipo": "reconocido",
                "mensaje": "Se reconoció a Estudiante de Prueba.",
            },
        )

        self.assertEqual(response.status_code, 404)

    def test_phone_cannot_start_a_tablet_that_is_not_online(self):
        response = self.client.post(
            "/api/camara/tablet/comando",
            json={
                "tablet_id": "tablet-offline",
                "action": "iniciar",
                "materia_id": 1,
            },
        )

        self.assertEqual(response.status_code, 404)

    def test_tablet_can_unregister_itself(self):
        self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-001",
                "nombre": "Tablet AsisTech",
                "activa": False,
            },
        )
        response = self.client.post(
            "/api/camara/tablet/heartbeat",
            json={
                "tablet_id": "tablet-001",
                "nombre": "Tablet AsisTech",
                "activa": False,
                "habilitada": False,
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            self.client.get("/api/camara/tablet/dispositivos").get_json(),
            [],
        )


if __name__ == "__main__":
    unittest.main()
